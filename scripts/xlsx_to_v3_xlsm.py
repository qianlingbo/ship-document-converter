#!/usr/bin/env python3
"""
xlsx_to_v3_xlsm.py — 将v2 xlsx数据迁移到v3 xlsm模板

用法:
  python3 xlsx_to_v3_xlsm.py <源xlsx路径> <船名> <日期> [输出目录]

示例:
  python3 xlsx_to_v3_xlsm.py 单证录入_COSCO_20260617.xlsx "COSCO SHIPPING WISDOM" 20260617
  python3 xlsx_to_v3_xlsm.py 单证录入_KANGSHUN99_20260620.xlsx KANGSHUN99 20260620

输出:
  oneTableDeclaration_{船名}_{日期}.xlsm
"""
import openpyxl, shutil, sys
from pathlib import Path
from copy import copy

TPL = Path(__file__).parent.parent / "templates" / "单证录入标准格式_v3.xlsm"

def copy_cell_style(src_cell, dst_cell):
    if not src_cell.has_style:
        return
    try:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
    except Exception:
        pass

def convert(xlsx_path, ship, date, out_dir=None):
    xlsx_path = Path(xlsx_path)
    if out_dir is None:
        out_dir = xlsx_path.parent
    else:
        out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"oneTableDeclaration_{ship}_{date}.xlsm"

    shutil.copy(TPL, out_file)

    wb_src = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_dst = openpyxl.load_workbook(out_file, keep_vba=True)

    # v2 xlsx: 数据从Row2起; v3 xlsm: 数据从Row3起
    SRC_START, DST_START = 2, 3

    for sn in ["船上非旅客人员清单", "船上非旅客人员物品清单", "海事船岸活动信息"]:
        ws_src = wb_src[sn]
        ws_dst = wb_dst[sn]
        for r in range(SRC_START, ws_src.max_row + 1):
            row_vals = [ws_src.cell(r, c).value
                        for c in range(1, (ws_src.max_column or 16) + 1)]
            if not any(v is not None for v in row_vals):
                continue
            target_r = r - SRC_START + DST_START
            for ci, val in enumerate(row_vals):
                src_c = ws_src.cell(r, ci + 1)
                dst_c = ws_dst.cell(target_r, ci + 1)
                dst_c.value = val
                copy_cell_style(src_c, dst_c)

    wb_dst.save(out_file)
    print(f"生成: {out_file}")
    return out_file

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    xlsx_path, ship, date = sys.argv[1], sys.argv[2], sys.argv[3]
    out_dir = sys.argv[4] if len(sys.argv) > 4 else None
    convert(xlsx_path, ship, date, out_dir)
